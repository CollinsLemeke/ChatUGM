# --- UGM Chatbot for Hugging Face Spaces ---
# Install dependencies in requirements.txt

import os
import json
import logging
import random
import uuid
import csv
from datetime import datetime
import torch
from transformers import pipeline, AutoTokenizer
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import gradio as gr
import pandas as pd
import re

# Document processing imports for Study Mode
try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

# --- Logging ---
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("UGM Chatbot")

# --- Config (edit these paths/names as needed) ---
HF_TOKEN = os.environ.get("HF_TOKEN", "")
MODEL_ID = "meta-llama/Llama-3.2-1B-Instruct"
DEVICE_MAP = "auto"
TORCH_DTYPE = torch.float16

# Path to the refined intents JSON (place in same directory or use HF datasets)
INTENTS_PATH = "intents.json"

# Where chat logs will be stored
LOG_DIR = "chat_logs"
os.makedirs(LOG_DIR, exist_ok=True)
MASTER_LOG_CSV = os.path.join(LOG_DIR, "chat_history.csv")
QUESTIONS_CSV = os.path.join(LOG_DIR, "questions.csv")

if not os.path.exists(QUESTIONS_CSV):
    with open(QUESTIONS_CSV, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["timestamp_utc","session_id","user_text","intent_tag","score"])


def preprocess_text(text):
    if not isinstance(text, str):
        return ""
    text = text.strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def load_intents(path=INTENTS_PATH):
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data.get("intents", [])
    except:
        return []


intents = load_intents()

rows = []
for intent in intents:
    tag = intent.get("tag")
    responses = intent.get("responses", [])
    if isinstance(responses, list) and responses:
        response = random.choice(responses).strip()
    else:
        response = intent.get("response", "").strip()

    for p in intent.get("patterns", []):
        rows.append({
            "pattern": p,
            "pattern_norm": preprocess_text(p),
            "tag": tag,
            "response": response
        })

corpus = [r["pattern_norm"] for r in rows]
tfidf = TfidfVectorizer(ngram_range=(1,2), lowercase=False).fit(corpus) if corpus else None
X = tfidf.transform(corpus) if tfidf is not None else None


def match_intent(text, threshold=0.45):
    if tfidf is None:
        return None, 0.0
    text_norm = preprocess_text(text)
    vec = tfidf.transform([text_norm])
    sims = cosine_similarity(vec, X)[0]
    idx = sims.argmax()
    score = float(sims[idx])
    if score >= threshold:
        return rows[idx], score
    return None, float(score)


# =====================================================
# STUDY MODE - Document Processing & Interactive Q&A
# =====================================================
class StudyMode:
    def __init__(self):
        self.document_text = None
        self.chunks = []
        self.interact_active = False
        self.assignment_active = False
        self.conversation_history = []
        self.assignment_breakdown = None
    
    def extract_text(self, file_obj):
        """Extract text from uploaded document"""
        if file_obj is None:
            return None, "No file uploaded."
        
        # Handle both file path string and file object
        if hasattr(file_obj, 'name'):
            file_path = file_obj.name
        else:
            file_path = str(file_obj)
        
        if not os.path.exists(file_path):
            return None, f"File not found: {file_path}"
        
        ext = os.path.splitext(file_path)[1].lower()
        
        try:
            if ext == '.pdf':
                return self._extract_pdf(file_path)
            elif ext in ['.docx', '.doc']:
                return self._extract_docx(file_path)
            elif ext in ['.pptx', '.ppt']:
                return self._extract_pptx(file_path)
            elif ext in ['.xlsx', '.xls']:
                return self._extract_excel(file_path)
            elif ext == '.csv':
                return self._extract_csv(file_path)
            elif ext == '.json':
                return self._extract_json(file_path)
            elif ext == '.txt':
                return self._extract_txt(file_path)
            else:
                return None, f"Unsupported file type: {ext}"
        except Exception as e:
            logger.error(f"Error extracting text: {e}")
            return None, f"Error reading file: {str(e)}"
    
    def _extract_pdf(self, file_path):
        if pdfplumber is None:
            return None, "PDF support not available. Install pdfplumber."
        try:
            text = []
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        text.append(t)
            if text:
                return "\n\n".join(text), None
            return None, "Could not extract text from PDF."
        except Exception as e:
            return None, f"PDF error: {str(e)}"
    
    def _extract_docx(self, file_path):
        if DocxDocument is None:
            return None, "Word support not available. Install python-docx."
        try:
            doc = DocxDocument(file_path)
            text = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also get text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        text.append(row_text)
            if text:
                return "\n\n".join(text), None
            return None, "Could not extract text from Word document."
        except Exception as e:
            return None, f"Word error: {str(e)}"
    
    def _extract_pptx(self, file_path):
        if Presentation is None:
            return None, "PowerPoint support not available. Install python-pptx."
        try:
            prs = Presentation(file_path)
            text = []
            for slide_num, slide in enumerate(prs.slides, 1):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                if slide_text:
                    text.append(f"[Slide {slide_num}]\n" + "\n".join(slide_text))
            if text:
                return "\n\n".join(text), None
            return None, "Could not extract text from PowerPoint."
        except Exception as e:
            return None, f"PowerPoint error: {str(e)}"
    
    def _extract_excel(self, file_path):
        try:
            df_dict = pd.read_excel(file_path, sheet_name=None)
            text = []
            for sheet, data in df_dict.items():
                text.append(f"[Sheet: {sheet}]")
                # Add headers
                text.append(" | ".join(str(col) for col in data.columns))
                for _, row in data.iterrows():
                    row_text = " | ".join(str(v) for v in row.values if pd.notna(v))
                    if row_text.strip():
                        text.append(row_text)
            if text:
                return "\n".join(text), None
            return None, "Could not extract text from Excel."
        except Exception as e:
            return None, f"Excel error: {str(e)}"
    
    def _extract_csv(self, file_path):
        try:
            df = pd.read_csv(file_path)
            text = [" | ".join(str(col) for col in df.columns)]
            for _, row in df.iterrows():
                row_text = " | ".join(str(v) for v in row.values if pd.notna(v))
                if row_text.strip():
                    text.append(row_text)
            if text:
                return "\n".join(text), None
            return None, "Could not extract text from CSV."
        except Exception as e:
            return None, f"CSV error: {str(e)}"
    
    def _extract_json(self, file_path):
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
            
            def flatten_json(obj, prefix=""):
                items = []
                if isinstance(obj, dict):
                    for k, v in obj.items():
                        items.extend(flatten_json(v, f"{k}: "))
                elif isinstance(obj, list):
                    for i, item in enumerate(obj):
                        items.extend(flatten_json(item, f"Item {i+1}: "))
                else:
                    if obj is not None and str(obj).strip():
                        items.append(f"{prefix}{obj}")
                return items
            
            text = flatten_json(data)
            if text:
                return "\n".join(text), None
            return None, "Could not extract text from JSON."
        except Exception as e:
            return None, f"JSON error: {str(e)}"
    
    def _extract_txt(self, file_path):
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            if content.strip():
                return content, None
            return None, "Text file is empty."
        except Exception as e:
            return None, f"Text file error: {str(e)}"
    
    def process_document(self, file_obj):
        """Process uploaded document"""
        text, error = self.extract_text(file_obj)
        
        if error:
            return False, f"❌ {error}"
        
        if not text or len(text.strip()) < 50:
            return False, "❌ Document is empty or too short."
        
        self.document_text = text
        self.chunks = self._create_chunks(text)
        self.interact_active = False
        self.assignment_active = False
        self.assignment_breakdown = None
        self.conversation_history = []
        
        if len(self.chunks) < 1:
            return False, "❌ Document too short. Need more content."
        
        return True, f"✅ Document loaded successfully! Found {len(self.chunks)} sections.\n\n**Choose your study mode:**\n• Type **interact** to ask questions about your document\n• Type **assignment** to get help with understanding and completing assignments\n• Type **summary** to get a comprehensive summary with key points"
    
    def _create_chunks(self, text, min_size=100, max_size=500):
        """Split text into meaningful chunks for better context retrieval"""
        # Split by double newlines or paragraph markers
        paragraphs = re.split(r'\n\s*\n|\n(?=[A-Z\[])', text)
        chunks = []
        current = ""
        
        for p in paragraphs:
            p = p.strip()
            if not p or len(p) < 30:
                continue
            if len(current) + len(p) > max_size and len(current) >= min_size:
                chunks.append(current.strip())
                current = p
            else:
                current = (current + " " + p).strip() if current else p
        
        if current and len(current) >= min_size // 2:
            chunks.append(current.strip())
        
        # If still no good chunks, split by sentences
        if len(chunks) < 1:
            sentences = re.split(r'(?<=[.!?])\s+', text)
            chunks = []
            current = ""
            for s in sentences:
                s = s.strip()
                if not s:
                    continue
                if len(current) + len(s) > max_size and current:
                    chunks.append(current.strip())
                    current = s
                else:
                    current = (current + " " + s).strip() if current else s
            if current:
                chunks.append(current.strip())
        
        return chunks
    
    def _find_relevant_chunks(self, query, top_k=3):
        """Find the most relevant chunks for a given query using TF-IDF similarity"""
        if not self.chunks:
            return []
        
        # Create TF-IDF vectorizer for chunks
        chunk_vectorizer = TfidfVectorizer(ngram_range=(1, 2), lowercase=True, stop_words='english')
        
        try:
            chunk_vectors = chunk_vectorizer.fit_transform(self.chunks)
            query_vector = chunk_vectorizer.transform([query.lower()])
            
            # Calculate cosine similarity
            similarities = cosine_similarity(query_vector, chunk_vectors)[0]
            
            # Get top k most similar chunks
            top_indices = similarities.argsort()[-top_k:][::-1]
            
            relevant_chunks = []
            for idx in top_indices:
                if similarities[idx] > 0.05:  # Minimum similarity threshold
                    relevant_chunks.append({
                        'text': self.chunks[idx],
                        'score': similarities[idx]
                    })
            
            return relevant_chunks
        except Exception as e:
            logger.error(f"Error finding relevant chunks: {e}")
            # Fallback: return first few chunks
            return [{'text': chunk, 'score': 0.5} for chunk in self.chunks[:top_k]]
    
    def generate_summary(self, pipe=None, tokenizer=None):
        """Generate a comprehensive summary with key points"""
        if not self.document_text:
            return "No document loaded. Please upload a study material first."
        
        # Try LLM-based summary first
        if pipe and tokenizer:
            result = self._llm_summary(pipe, tokenizer)
            if result:
                return result
        
        # Fallback to extractive summary
        return self._extractive_summary()
    
    def _llm_summary(self, pipe, tokenizer):
        """Use LLM to generate summary"""
        # Use first ~1500 chars to fit in context
        doc_excerpt = self.document_text[:1500]
        
        prompt = f"""Summarize the following document with key points. Be comprehensive and organized.

Document:
{doc_excerpt}

Provide a summary with:
1. A brief overview (2-3 sentences)
2. Key points (bullet format)
3. Main takeaways

Summary:"""
        
        try:
            out = pipe(prompt, max_new_tokens=400, do_sample=True, temperature=0.7,
                      pad_token_id=tokenizer.eos_token_id if tokenizer else 50256)
            gen = out[0].get("generated_text", "")
            if gen.startswith(prompt):
                gen = gen[len(prompt):]
            
            summary = gen.strip()
            if len(summary) > 50:
                return f"📄 **Document Summary**\n\n{summary}"
            return None
        except Exception as e:
            logger.error(f"LLM summary error: {e}")
            return None
    
    def _extractive_summary(self):
        """Fallback extractive summary using key sentences"""
        text = self.document_text
        
        # Split into sentences
        sentences = re.split(r'(?<=[.!?])\s+', text)
        sentences = [s.strip() for s in sentences if len(s.strip()) > 30]
        
        if not sentences:
            return "📄 **Document Summary**\n\nDocument content is too short to summarize."
        
        # Extract key information
        summary_parts = []
        
        # 1. Overview - first few meaningful sentences
        overview_sentences = sentences[:3]
        overview = " ".join(overview_sentences)
        summary_parts.append(f"**📋 Overview:**\n{overview}")
        
        # 2. Key Points - extract sentences with important markers
        key_markers = ['important', 'key', 'main', 'significant', 'essential', 'critical', 
                       'must', 'should', 'note', 'remember', 'conclusion', 'result']
        key_sentences = []
        for sent in sentences:
            sent_lower = sent.lower()
            if any(marker in sent_lower for marker in key_markers):
                key_sentences.append(sent)
            # Also include sentences with numbers/statistics
            elif re.search(r'\d+%|\d+\.\d+|\$\d+', sent):
                key_sentences.append(sent)
        
        # Get unique key points (max 5)
        key_sentences = list(dict.fromkeys(key_sentences))[:5]
        
        if key_sentences:
            key_points = "\n".join([f"• {s}" for s in key_sentences])
            summary_parts.append(f"\n**🔑 Key Points:**\n{key_points}")
        
        # 3. Extract names, terms, and concepts
        concepts = re.findall(r'\b[A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+)*\b', text)
        concepts = [c for c in concepts if len(c) > 3 and c not in ['The', 'This', 'That', 'These', 'Those', 'When', 'Where', 'What', 'Which', 'There']]
        unique_concepts = list(dict.fromkeys(concepts))[:8]
        
        if unique_concepts:
            summary_parts.append(f"\n**📌 Key Terms & Concepts:**\n{', '.join(unique_concepts)}")
        
        # 4. Statistics/Numbers mentioned
        numbers = re.findall(r'(?:\$?\d+(?:,\d{3})*(?:\.\d+)?(?:%|\s*percent)?)', text)
        unique_numbers = list(dict.fromkeys(numbers))[:5]
        if unique_numbers:
            summary_parts.append(f"\n**📊 Key Figures:**\n{', '.join(unique_numbers)}")
        
        # 5. Document stats
        word_count = len(text.split())
        section_count = len(self.chunks)
        summary_parts.append(f"\n**📈 Document Statistics:**\n• Word count: ~{word_count}\n• Sections: {section_count}")
        
        # 6. Main takeaway (last meaningful sentence or conclusion)
        conclusion_markers = ['conclusion', 'summary', 'finally', 'in conclusion', 'to summarize', 'overall']
        takeaway = None
        for sent in reversed(sentences):
            if any(marker in sent.lower() for marker in conclusion_markers):
                takeaway = sent
                break
        if not takeaway and len(sentences) > 3:
            takeaway = sentences[-1]
        
        if takeaway:
            summary_parts.append(f"\n**💡 Main Takeaway:**\n{takeaway}")
        
        return "📄 **Document Summary**\n\n" + "\n".join(summary_parts) + "\n\n---\nType **interact** to ask questions, **assignment** for help, or **summary** to see this again."
    
    def start_interact(self):
        """Start interactive mode"""
        self.interact_active = True
        self.assignment_active = False
        self.conversation_history = []
        return "💬 **Interactive Mode Active!**\n\nI'm ready to answer your questions about the document. Just ask me anything and I'll find the relevant information for you.\n\nType **exit** to leave interactive mode, **assignment** for assignment help, or **summary** to see the document summary."
    
    def start_assignment_mode(self, pipe=None, tokenizer=None):
        """Start assignment mode with breakdown"""
        self.assignment_active = True
        self.interact_active = False
        self.conversation_history = []
        
        if not self.document_text:
            return "No assignment document loaded. Please upload your assignment first."
        
        # Generate assignment breakdown
        breakdown = self._analyze_assignment(pipe, tokenizer)
        self.assignment_breakdown = breakdown
        
        return breakdown
    
    def _analyze_assignment(self, pipe=None, tokenizer=None):
        """Analyze and breakdown the assignment document"""
        text = self.document_text
        
        # Try LLM-based analysis first
        if pipe and tokenizer:
            result = self._llm_assignment_breakdown(pipe, tokenizer)
            if result:
                return result
        
        # Fallback to rule-based analysis
        return self._extractive_assignment_breakdown()
    
    def _llm_assignment_breakdown(self, pipe, tokenizer):
        """Use LLM to generate assignment breakdown"""
        doc_excerpt = self.document_text[:1500]
        
        prompt = f"""You are a helpful study assistant. Analyze this assignment document and provide a clear breakdown to help the student understand and complete it.

Assignment Document:
{doc_excerpt}

Provide:
1. **Assignment Overview**: What is this assignment about? (2-3 sentences)
2. **Main Tasks**: List the key tasks or questions that need to be completed
3. **Requirements**: What are the specific requirements or guidelines?
4. **Key Concepts**: What concepts or topics should the student focus on?
5. **Suggested Approach**: Step-by-step guidance on how to tackle this assignment

Be clear, helpful, and encouraging. Break down complex requirements into simple steps.

Breakdown:"""
        
        try:
            out = pipe(prompt, max_new_tokens=500, do_sample=True, temperature=0.7,
                      pad_token_id=tokenizer.eos_token_id if tokenizer else 50256)
            gen = out[0].get("generated_text", "")
            if gen.startswith(prompt):
                gen = gen[len(prompt):]
            
            response = gen.strip()
            if len(response) > 50:
                return f"📝 **Assignment Analysis**\n\n{response}\n\n---\n💡 Ask me questions about specific parts of the assignment!\nType **exit** to leave assignment mode."
            return None
        except Exception as e:
            logger.error(f"LLM assignment breakdown error: {e}")
            return None
    
    def _extractive_assignment_breakdown(self):
        """Fallback rule-based assignment breakdown"""
        text = self.document_text
        
        # Identify assignment parts
        parts = []
        
        # 1. Extract title/assignment name
        first_lines = text.split('\n')[:5]
        title = None
        for line in first_lines:
            if len(line.strip()) > 10 and len(line.strip()) < 150:
                title = line.strip()
                break
        
        if title:
            parts.append(f"**📋 Assignment Title:**\n{title}\n")
        
        # 2. Find questions/tasks (numbered items, bullets, or "Question" markers)
        question_patterns = [
            r'(?:Question|Task|Part|Section|Problem)\s*\d+[:\.]?\s*(.+?)(?=(?:Question|Task|Part|Section|Problem)\s*\d+|$)',
            r'^\d+\.\s*(.+?)(?=^\d+\.|$)',
            r'^[a-z]\)\s*(.+?)(?=^[a-z]\)|$)',
        ]
        
        tasks = []
        for pattern in question_patterns:
            matches = re.findall(pattern, text, re.MULTILINE | re.DOTALL | re.IGNORECASE)
            if matches:
                tasks.extend([m.strip()[:200] for m in matches if len(m.strip()) > 20])
                break
        
        if tasks:
            tasks = list(dict.fromkeys(tasks))[:8]  # Unique, max 8
            tasks_text = "\n".join([f"• {task}" for task in tasks])
            parts.append(f"**📌 Main Tasks:**\n{tasks_text}\n")
        
        # 3. Extract requirements (look for words like "must", "should", "required")
        requirement_markers = ['must', 'should', 'required', 'need to', 'have to', 'ensure', 'make sure']
        sentences = re.split(r'(?<=[.!?])\s+', text)
        requirements = []
        for sent in sentences:
            sent_lower = sent.lower()
            if any(marker in sent_lower for marker in requirement_markers) and len(sent) > 30:
                requirements.append(sent.strip())
        
        if requirements:
            requirements = list(dict.fromkeys(requirements))[:5]
            req_text = "\n".join([f"• {req}" for req in requirements])
            parts.append(f"**✅ Requirements:**\n{req_text}\n")
        
        # 4. Extract deadlines or dates
        dates = re.findall(r'\b(?:due|deadline|submit(?:ted)?|by)\s*:?\s*([A-Za-z]+\s+\d{1,2}(?:st|nd|rd|th)?(?:,?\s+\d{4})?|\d{1,2}[/-]\d{1,2}[/-]\d{2,4})', text, re.IGNORECASE)
        if dates:
            dates = list(dict.fromkeys(dates))[:3]
            parts.append(f"**📅 Important Dates:**\n{', '.join(dates)}\n")
        
        # 5. Extract key concepts/topics
        concepts = re.findall(r'\b[A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+)*\b', text)
        concepts = [c for c in concepts if len(c) > 3 and c not in ['The', 'This', 'That', 'Assignment', 'Question', 'Task']]
        unique_concepts = list(dict.fromkeys(concepts))[:10]
        
        if unique_concepts:
            parts.append(f"**🎯 Key Topics:**\n{', '.join(unique_concepts)}\n")
        
        # 6. Suggested approach
        approach = """**💡 Suggested Approach:**
1. Read through the entire assignment carefully
2. Identify what's being asked in each task
3. Break down complex questions into smaller parts
4. Research relevant concepts and gather information
5. Plan your answers before writing
6. Review requirements and ensure you've addressed everything
7. Proofread and check your work before submission"""
        
        parts.append(approach)
        
        if not parts:
            return "📝 **Assignment Analysis**\n\nI've loaded your assignment document. Ask me questions about specific parts and I'll help you understand and complete them!\n\nType **exit** to leave assignment mode."
        
        return "📝 **Assignment Analysis**\n\n" + "\n".join(parts) + "\n\n---\n💡 Ask me questions about specific parts of the assignment!\nType **exit** to leave assignment mode, **interact** for general questions, or **summary** for document overview."
    
    def handle_assignment_question(self, user_query, pipe=None, tokenizer=None):
        """Handle questions in assignment mode"""
        if not self.assignment_active:
            return "Type **assignment** first to start assignment mode."
        
        if not self.document_text:
            return "No assignment loaded. Please upload your assignment document first."
        
        # Check for exit command
        if user_query.lower().strip() in ['exit', 'quit', 'stop', 'end', 'done']:
            self.assignment_active = False
            return "👋 Left assignment mode. Type **assignment** to start again, **interact** for questions, or **summary** for overview."
        
        # Find relevant content
        relevant_chunks = self._find_relevant_chunks(user_query, top_k=3)
        
        if not relevant_chunks:
            return "I couldn't find specific information about that in the assignment. Could you rephrase or ask about a specific part? 🤔"
        
        # Build context
        context = "\n\n".join([chunk['text'] for chunk in relevant_chunks])
        
        # Try LLM-based response
        if pipe and tokenizer:
            response = self._llm_assignment_response(user_query, context, pipe, tokenizer)
            if response:
                return response
        
        # Fallback response
        return self._fallback_assignment_response(user_query, relevant_chunks)
    
    def _llm_assignment_response(self, query, context, pipe, tokenizer):
        """Generate LLM-based response for assignment questions"""
        prompt = f"""You are a helpful tutor guiding a student through their assignment. Help them understand the assignment requirements and guide them toward completing it, but don't do the work for them.

Assignment Context:
{context[:1200]}

Student Question: {query}

Instructions:
- Help the student understand what's being asked
- Break down complex requirements into simple steps
- Guide them on how to approach the task
- Encourage critical thinking
- Don't provide direct answers, but help them think through the problem
- Be supportive and encouraging

Response:"""
        
        try:
            out = pipe(prompt, max_new_tokens=300, do_sample=True, temperature=0.7,
                      pad_token_id=tokenizer.eos_token_id if tokenizer else 50256)
            gen = out[0].get("generated_text", "")
            if gen.startswith(prompt):
                gen = gen[len(prompt):]
            
            response = gen.strip()
            if len(response) > 20:
                return response + "\n\n💡 Any other questions about the assignment?"
            return None
        except Exception as e:
            logger.error(f"LLM assignment response error: {e}")
            return None
    
    def _fallback_assignment_response(self, query, relevant_chunks):
        """Fallback response for assignment questions"""
        if not relevant_chunks:
            return "I couldn't find information about that in your assignment. Try asking about a specific task or requirement? 🤔"
        
        # Present relevant content
        best_chunk = relevant_chunks[0]
        text = best_chunk['text']
        
        # Extract key sentences
        sentences = re.split(r'(?<=[.!?])\s+', text)
        
        # Look for question/task markers
        task_sentences = []
        for sent in sentences:
            if any(marker in sent.lower() for marker in ['question', 'task', 'write', 'explain', 'describe', 'analyze', 'discuss', 'must', 'should']):
                task_sentences.append(sent)
        
        if task_sentences:
            response = "**Here's what the assignment says:**\n\n"
            response += "\n\n".join([f"• {sent}" for sent in task_sentences[:3]])
            response += "\n\n**To approach this:**\n"
            response += "1. Break down what's being asked\n"
            response += "2. Think about the key concepts involved\n"
            response += "3. Plan your answer structure\n"
            response += "4. Gather relevant information\n"
            response += "5. Write clearly and check your work"
        else:
            snippet = text[:400] + "..." if len(text) > 400 else text
            response = f"**Here's the relevant part from your assignment:**\n\n\"{snippet}\"\n\n"
            response += "**Need help with:**\n"
            response += "• Understanding what's required?\n"
            response += "• How to approach this task?\n"
            response += "• Key concepts to focus on?\n\n"
            response += "Just ask! 📚"
        
        return response
    
    def handle_interact(self, user_query, pipe=None, tokenizer=None):
        """Handle user questions about the document"""
        if not self.interact_active:
            return "Type **interact** first to start asking questions about your document."
        
        if not self.document_text:
            return "No document loaded. Please upload a study material first."
        
        # Check for exit command
        if user_query.lower().strip() in ['exit', 'quit', 'stop', 'end', 'done']:
            self.interact_active = False
            return "👋 Left interactive mode. Type **interact** to start again, **assignment** for help, or **summary** for document overview."
        
        # Find relevant chunks
        relevant_chunks = self._find_relevant_chunks(user_query)
        
        if not relevant_chunks:
            return "I couldn't find specific information about that in the document. Could you try rephrasing your question? 🤔"
        
        # Build context from relevant chunks
        context = "\n\n".join([chunk['text'] for chunk in relevant_chunks])
        
        # Try LLM-based response
        if pipe and tokenizer:
            response = self._llm_interact_response(user_query, context, pipe, tokenizer)
            if response:
                return response
        
        # Fallback to direct context presentation
        return self._fallback_interact_response(user_query, relevant_chunks)
    
    def _llm_interact_response(self, query, context, pipe, tokenizer):
        """Generate LLM-based response to user query"""
        prompt = f"""You are a helpful study assistant. Answer the user's question based ONLY on the document context provided. Be friendly, clear, and cite specific information from the document.

Document Context:
{context[:1200]}

User Question: {query}

Instructions:
- Answer based only on the document content
- Be conversational and helpful
- If the answer isn't in the document, say so politely
- Keep the response concise but complete

Answer:"""
        
        try:
            out = pipe(prompt, max_new_tokens=300, do_sample=True, temperature=0.7,
                      pad_token_id=tokenizer.eos_token_id if tokenizer else 50256)
            gen = out[0].get("generated_text", "")
            if gen.startswith(prompt):
                gen = gen[len(prompt):]
            
            response = gen.strip()
            if len(response) > 20:
                return response
            return None
        except Exception as e:
            logger.error(f"LLM interact error: {e}")
            return None
    
    def _fallback_interact_response(self, query, relevant_chunks):
        """Fallback response when LLM is unavailable"""
        if not relevant_chunks:
            return "I couldn't find information about that in your document. Try asking something else? 🤔"
        
        # Present the most relevant information
        best_chunk = relevant_chunks[0]
        text = best_chunk['text']
        
        # Try to extract the most relevant sentence
        sentences = re.split(r'(?<=[.!?])\s+', text)
        query_words = set(query.lower().split())
        
        best_sentence = None
        best_overlap = 0
        
        for sent in sentences:
            sent_words = set(sent.lower().split())
            overlap = len(query_words & sent_words)
            if overlap > best_overlap:
                best_overlap = overlap
                best_sentence = sent
        
        if best_sentence and len(best_sentence) > 30:
            response = f"Based on the document:\n\n\"{best_sentence}\""
        else:
            # Show a snippet of the relevant chunk
            snippet = text[:300] + "..." if len(text) > 300 else text
            response = f"Here's what I found in the document:\n\n\"{snippet}\""
        
        response += "\n\nFeel free to ask more questions! 📚"
        return response
    
    def reset(self):
        """Reset study mode"""
        self.document_text = None
        self.chunks = []
        self.interact_active = False
        self.assignment_active = False
        self.assignment_breakdown = None
        self.conversation_history = []


# --- Chatbot class ---
class UGMChatbot:
    def __init__(self):
        self.system_msg = "You are Axiom, The University of Greater Manchester's helpful assistant."
        self.messages = [{"role":"system","content": self.system_msg}]
        self.tokenizer = None
        self.pipe = None
        self.study_mode = StudyMode()
        self.setup_model()
        if not os.path.exists(MASTER_LOG_CSV):
            with open(MASTER_LOG_CSV, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(["timestamp","session_id","user_text","assistant_text","intent_tag","score"])
        logger.info("UGMChatbot initialized.")

    def setup_model(self):
        try:
            if HF_TOKEN:
                from huggingface_hub import login
                login(token=HF_TOKEN)
            self.tokenizer = AutoTokenizer.from_pretrained(MODEL_ID, token=HF_TOKEN if HF_TOKEN else None)
            self.pipe = pipeline(
                "text-generation",
                model=MODEL_ID,
                tokenizer=self.tokenizer,
                device_map=DEVICE_MAP,
                torch_dtype=TORCH_DTYPE,
            )
            logger.info("Model loaded (device_map=%s).", DEVICE_MAP)
        except Exception as e:
            logger.warning("GPU/model load failed (%s). Trying CPU...", e)
            try:
                self.pipe = pipeline("text-generation", model=MODEL_ID, device=-1)
                self.tokenizer = AutoTokenizer.from_pretrained(MODEL_ID)
            except Exception as e2:
                logger.error("Fallback failed: %s", e2)
                self.pipe = None
                self.tokenizer = None

    def save_interaction(self, session_id, user_text, assistant_text, tag, score):
        row = [datetime.utcnow().isoformat(), session_id, user_text, assistant_text, tag or "", score or 0.0]
        with open(MASTER_LOG_CSV, "a", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(row)
        qrow = [datetime.utcnow().isoformat(), session_id, user_text, tag or "", score or 0.0]
        with open(QUESTIONS_CSV, "a", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(qrow)

    def generate_response(self, user_text, session_id):
        matched, score = match_intent(user_text)
        if matched:
            tag = matched.get("tag")
            reply = matched.get("response", "Sorry, I don't have an answer for that at the moment - I am still learning 🙂.")
        else:
            tag = None
            if self.pipe is None:
                reply = "Sorry — Please be specific on your questions or rather type one-word like courses, fees, etc 🙂."
            else:
                self.messages.append({"role":"user","content": user_text})
                try:
                    prompt_parts = [m["content"] for m in self.messages[-6:]]
                    prompt = "\n".join(prompt_parts) + "\nAssistant:"
                    outputs = self.pipe(
                        prompt,
                        max_new_tokens=200,
                        do_sample=True,
                        temperature=0.7,
                        top_p=0.9,
                        pad_token_id=self.tokenizer.eos_token_id if self.tokenizer else 50256,
                    )
                    generated = outputs[0].get("generated_text", "")
                    if generated.startswith(prompt):
                        generated = generated[len(prompt):]
                    reply = generated.strip()
                    self.messages.append({"role":"assistant","content": reply})
                except Exception as e:
                    logger.error("Model generation error: %s", e)
                    reply = "Sorry — model generation failed. Try rephrasing your question 🙂."
        self.save_interaction(session_id, user_text, reply, tag, score)
        return reply, (tag or ""), float(score or 0.0)


# --- Gradio Interface ---
def ensure_session_id(state):
    if state is None or not isinstance(state, dict) or "session_id" not in state:
        return {"session_id": str(uuid.uuid4()), "history": []}
    return state

bot = UGMChatbot()

# ---------- CSS ----------
css = """
*, *::before, *::after { box-sizing: border-box; }
html, body {
  margin: 0; padding: 0; width: 100%; height: 100%;
  background: #0f0f10 !important; color: #eee !important;
  overflow-x: hidden; color-scheme: dark !important;
}

:root, .app.svelte-4a45dv.svelte-4a45dv, .block, .gradio-container {
  width: 100vw !important; max-width: 100vw !important;
  margin: 0 !important; padding: 0 !important;
  box-sizing: border-box; overflow-x: hidden !important;
}

.gradio-container, .main, .app, .contain, .gap, .form,
.block, .wrap, .panel, [data-testid="chatbot"] {
  background: #0f0f10 !important;
  color: #eee !important;
  border-color: #2a2a2e !important;
}

#ugm-logo {
  display: flex !important; align-items: center !important; justify-content: center !important;
  margin: 0 !important; padding: 0 !important;
  height: 50px !important; width: 50px !important;
  min-width: 50px !important; max-width: 50px !important; min-height: 50px !important;
  background: transparent !important; border: none !important; overflow: hidden !important;
}

#ugm-logo img {
  height: 46px !important; width: 46px !important;
  object-fit: contain !important; border-radius: 8px !important;
  margin: 0 !important; padding: 0 !important;
  display: block !important; background: transparent !important;
}

#ugm-logo .image-container, #ugm-logo .upload-container, #ugm-logo > div,
#ugm-logo .svelte-1p9xokt, #ugm-logo .image-frame {
  padding: 0 !important; margin: 0 !important; border: none !important;
  background: transparent !important;
  width: 50px !important; height: 50px !important;
  min-width: 50px !important; max-width: 50px !important;
  overflow: hidden !important; box-shadow: none !important;
}

#ugm-logo * { background: transparent !important; border: none !important; box-shadow: none !important; }

.header-row {
  position: fixed; top: 0; left: 0; width: 100vw; height: 64px;
  background: linear-gradient(90deg, #0058b0 0%, #d7263d 100%);
  display: flex !important; align-items: center !important; flex-direction: row !important;
  padding-left: 12px; padding-right: 120px;
  color: #fff; z-index: 50; gap: 0;
}

.header-row > div, .header-row > div > div { 
  display: flex !important; 
  align-items: center !important; 
  flex-direction: row !important;
}

.header-row .block, .header-row .prose, .header-row .markdown-text,
.header-row .svelte-1ed2p3z, .header-row [class*="svelte-"] {
  background: transparent !important;
  border: none !important;
  box-shadow: none !important;
  padding: 0 !important;
  margin: 0 !important;
  min-height: 0 !important;
  max-height: 64px !important;
}

.header-title { font-size: 1.2rem; font-weight: 600; color: #fff; white-space: nowrap; }

#reset-hat-btn {
  background: transparent !important;
  border: none !important;
  box-shadow: none !important;
  font-size: 32px !important;
  width: 44px !important;
  min-width: 44px !important;
  max-width: 44px !important;
  height: 44px !important;
  padding: 0 !important;
  margin: 0 4px 0 0 !important;
  cursor: pointer !important;
  display: flex !important;
  align-items: center !important;
  justify-content: center !important;
  transition: transform 0.2s ease !important;
}
#reset-hat-btn:hover { transform: scale(1.15) !important; }
#reset-hat-btn:active { transform: scale(0.95) !important; }

.chatbox {
  position: fixed !important; top: 64px !important;
  left: 0 !important; right: 0 !important; bottom: 80px !important;
  width: 100vw !important; max-width: 100vw !important;
  height: auto !important; overflow-y: auto !important; overflow-x: hidden !important;
  background: #0f0f10 !important;
  padding: 4px 4px !important; padding-bottom: 4px !important;
  border: none !important; border-radius: 0 !important;
  -webkit-overflow-scrolling: touch; transform: none !important;
  display: flex !important;
}

.chatbox > div { overflow-x: hidden !important; width: 100% !important; max-width: 100% !important; }

.chatbox .message {
  border-radius: 18px !important; padding: 10px 14px !important;
  margin: 6px 0 !important; word-wrap: break-word !important; overflow-wrap: break-word !important;
}

.chatbox .user, .chatbox [data-testid="user"], .chatbox .message-row.user-row .message-bubble {
  background: #003d80 !important; color: #ffffff !important;
  border-radius: 18px 18px 4px 18px !important;
  margin-left: auto !important; margin-right: 0 !important;
  max-width: 100% !important; width: auto !important;
  display: inline-block !important; white-space: normal !important;
  word-break: break-word !important; padding: 10px 14px 10px 12px !important;
  text-align: left !important; height: auto !important; min-height: unset !important;
}

.chatbox .user p, .chatbox .user span, .chatbox .user div,
.chatbox [data-testid="user"] p, .chatbox [data-testid="user"] span, .chatbox [data-testid="user"] div,
.chatbox .message-row.user-row .message-bubble p,
.chatbox .message-row.user-row .message-bubble span,
.chatbox .message-row.user-row .prose {
  margin: 0 !important; padding: 0 !important; text-align: left !important; white-space: normal !important;
  color: #ffffff !important;
}

.chatbox .bot, .chatbox .assistant, .chatbox [data-testid="bot"], .chatbox .message-row.bot-row .message-bubble {
  background: #1e1e22 !important; color: #ffffff !important;
  border-radius: 18px 18px 18px 4px !important;
  margin-right: auto !important; margin-left: 0 !important;
  max-width: 95% !important; width: auto !important;
  display: inline-block !important; white-space: normal !important;
  word-break: break-word !important; padding: 10px 14px 10px 12px !important; text-align: left !important;
}

.chatbox .bot p, .chatbox .bot span, .chatbox .bot div,
.chatbox .assistant p, .chatbox .assistant span, .chatbox .assistant div,
.chatbox [data-testid="bot"] p, .chatbox [data-testid="bot"] span, .chatbox [data-testid="bot"] div,
.chatbox .message-row.bot-row .message-bubble p,
.chatbox .message-row.bot-row .message-bubble span,
.chatbox .message-row.bot-row .prose,
.chatbox .message-row.bot-row .prose p {
  margin: 0 !important; padding: 0 !important; text-align: left !important;
  color: #ffffff !important;
}

.chatbox .message-row { display: flex !important; width: 100% !important; padding: 0 !important; margin: 4px 0 !important; }
.chatbox .message-row.user-row { justify-content: flex-end !important; }
.chatbox .message-row.bot-row { justify-content: flex-start !important; }
.chatbox .message-wrap { overflow-x: hidden !important; display: flex !important; flex-direction: column !important; padding: 0 !important; }
.chatbox .message-bubble-border { border: none !important; padding: 0 !important; }
.chatbox .message-bubble {
  border-radius: 18px !important; overflow-x: hidden !important;
  width: auto !important; max-width: 95% !important;
  display: inline-block !important; padding: 0 !important;
  height: auto !important; min-height: unset !important;
}
.chatbox .prose { padding: 0 !important; margin: 0 !important; }
.chatbox .prose p { margin: 0 !important; padding: 0 !important; }

/* Textbox Container - Fixed at very bottom of screen */
.textbox-container {
  position: fixed !important;
  inset: auto 0 0 0 !important;
  top: auto !important;
  bottom: 0 !important; 
  left: 0 !important; 
  right: 0 !important;
  transform: none !important;
  width: 100vw !important; 
  max-width: 100vw !important;
  background: #0f0f10 !important; 
  padding: 10px 20px 15px 20px !important; 
  z-index: 9999 !important;
  display: flex !important; 
  justify-content: center !important;
}

.textbox, .textbox textarea {
  width: 100% !important; font-size: 0.95rem !important;
  padding: 12px 56px 12px 18px !important;
  background: #252527 !important; color: #eee !important;
  border: 1px solid #3a3a3c !important; border-radius: 24px !important;
  outline: none !important; resize: none !important;
  overflow-y: hidden !important; overflow-x: hidden !important;
  min-height: 50px !important; max-height: 200px !important;
  line-height: 1.4 !important; box-shadow: 0 2px 10px rgba(0,0,0,0.4) !important;
  scrollbar-width: none !important; -ms-overflow-style: none !important;
}

.textbox textarea::-webkit-scrollbar { display: none !important; width: 0 !important; height: 0 !important; }
.textbox textarea::placeholder { color: #999 !important; font-size: 0.85rem !important; }
.textbox textarea:focus { border-color: #0058b0 !important; box-shadow: 0 0 0 2px rgba(0,88,176,0.2) !important; }

/* Send Button */
.send-btn {
  position: absolute !important; right: 36px !important; top: 50% !important;
  transform: translateY(-50%) !important;
  width: 32px !important; height: 32px !important;
  max-width: 32px !important; min-width: 32px !important;
  border: none !important;
  background: #0058b0 !important;
  color: #fff !important;
  font-size: 14px !important;
  cursor: pointer !important;
  outline: none !important;
  padding: 0 !important;
  border-radius: 50% !important;
  transition: background 0.2s ease, transform 0.1s ease !important;
  display: flex !important;
  align-items: center !important;
  justify-content: center !important;
  box-shadow: 0 2px 6px rgba(0,88,176,0.4) !important;
  z-index: 10 !important;
}
.send-btn:hover { background: #0070e0 !important; transform: translateY(-50%) scale(1.05) !important; }
.send-btn:active { background: #004a99 !important; transform: translateY(-50%) scale(0.95) !important; }
.send-btn::before { content: "➤" !important; line-height: 1 !important; }

/* STUDY MODE CONTROLS - TOP RIGHT */
.study-controls {
  position: fixed !important;
  top: 12px !important;
  right: 12px !important;
  bottom: auto !important;
  left: auto !important;
  display: flex !important;
  align-items: center !important;
  gap: 8px !important;
  height: 40px !important;
  z-index: 55 !important;
  background: transparent !important;
  padding: 0 !important;
  margin: 0 !important;
  border: none !important;
  flex-direction: row-reverse !important;
}

/* Study Mode Toggle Button */
#study-toggle-btn {
  width: 40px !important;
  height: 40px !important;
  min-width: 40px !important;
  max-width: 40px !important;
  min-height: 40px !important;
  max-height: 40px !important;
  border-radius: 50% !important;
  background: rgba(255,255,255,0.2) !important;
  border: none !important;
  color: #fff !important;
  font-size: 18px !important;
  cursor: pointer !important;
  display: flex !important;
  align-items: center !important;
  justify-content: center !important;
  transition: all 0.2s ease !important;
  padding: 0 !important;
  margin: 0 !important;
  box-shadow: 0 2px 8px rgba(0,0,0,0.3) !important;
}
#study-toggle-btn:hover {
  background: rgba(255,255,255,0.3) !important;
}
#study-toggle-btn.active {
  background: #28a745 !important;
  color: #fff !important;
}

/* FILE UPLOAD - CLEAN STYLING */
.study-file-upload {
  height: 36px !important;
  min-width: 36px !important;
  max-width: 36px !important;
  position: relative !important;
}

.study-file-upload > div {
  background: rgba(255,255,255,0.2) !important;
  border: none !important;
  border-radius: 50% !important;
  padding: 0 !important;
  height: 36px !important;
  width: 36px !important;
  display: flex !important;
  align-items: center !important;
  justify-content: center !important;
  cursor: pointer !important;
  transition: all 0.2s ease !important;
  box-shadow: 0 2px 8px rgba(0,0,0,0.3) !important;
  position: relative !important;
  overflow: hidden !important;
}

.study-file-upload > div:hover {
  background: rgba(255,255,255,0.3) !important;
}

.study-file-upload label,
.study-file-upload span,
.study-file-upload p,
.study-file-upload .file-preview,
.study-file-upload .upload-button,
.study-file-upload [data-testid="block-info"],
.study-file-upload .wrap,
.study-file-upload .or,
.study-file-upload .click-upload,
.study-file-upload svg {
  display: none !important;
  visibility: hidden !important;
}

.study-file-upload input[type="file"] { 
  opacity: 0 !important;
  position: absolute !important;
  width: 100% !important;
  height: 100% !important;
  cursor: pointer !important;
}

.study-file-upload .upload-area,
.study-file-upload [data-testid="upload-area"] {
  border: none !important;
  background: transparent !important;
  min-height: 36px !important;
  height: 36px !important;
  padding: 0 !important;
}

.study-file-upload::after {
  content: "📎" !important;
  position: absolute !important;
  top: 50% !important;
  left: 50% !important;
  transform: translate(-50%, -50%) !important;
  font-size: 18px !important;
  filter: brightness(0) invert(1) !important;
  pointer-events: none !important;
  z-index: 10 !important;
}

/* Hide copy and share buttons only */
.chatbox [aria-label="Copy"],
.chatbox [aria-label="Share"],
.chatbox button[title="Copy"],
.chatbox button[title="Share"],
.chatbox svg.feather-copy,
.chatbox svg.feather-share {
  display: none !important;
  visibility: hidden !important;
  opacity: 0 !important;
  pointer-events: none !important;
}

/* Scrollbar */
.chatbox::-webkit-scrollbar { width: 6px; }
.chatbox::-webkit-scrollbar-track { background: transparent; }
.chatbox::-webkit-scrollbar-thumb { background: #444; border-radius: 3px; }
.chatbox::-webkit-scrollbar-thumb:hover { background: #555; }
* { scrollbar-width: thin; }
*::-webkit-scrollbar-horizontal { display: none !important; height: 0 !important; }

/* MOBILE RESPONSIVE */
@media (max-width: 768px) {
  .chatbox { top: 64px !important; bottom: 76px !important; padding: 2px 2px !important; }
  .textbox-container {
    top: auto !important;
    bottom: 0px !important;
    width: calc(100vw - 20px) !important;
    max-width: calc(100vw - 20px) !important;
    left: 50% !important;
    transform: translateX(-50%) !important;
    z-index: 100 !important;
  }
  .textbox, .textbox textarea { font-size: 14px !important; min-height: 44px !important; padding: 10px 50px 10px 16px !important; }
  .textbox textarea::placeholder { font-size: 0.8rem !important; }
  
  .study-controls {
    top: 12px !important;
    right: 12px !important;
    height: 40px !important;
  }
  #study-toggle-btn {
    width: 36px !important;
    height: 36px !important;
    min-width: 36px !important;
    max-width: 36px !important;
    min-height: 36px !important;
    max-height: 36px !important;
    font-size: 16px !important;
  }
  .study-file-upload {
    min-width: 32px !important;
    max-width: 32px !important;
    height: 32px !important;
  }
  .study-file-upload > div {
    height: 32px !important;
    width: 32px !important;
  }
  
  .header-row { padding-left: 8px; padding-right: 110px !important; height: 60px !important; }
  #ugm-logo, #ugm-logo .image-container, #ugm-logo > div {
    width: 44px !important; height: 44px !important; min-width: 44px !important; max-width: 44px !important;
  }
  #ugm-logo img { width: 40px !important; height: 40px !important; }
}

@media (max-width: 480px) {
  .chatbox { top: 58px !important; bottom: 72px !important; padding: 2px 2px !important; }
  .textbox-container {
    top: auto !important;
    bottom: 0px !important;
    width: calc(100vw - 16px) !important;
    max-width: calc(100vw - 16px) !important;
    left: 50% !important;
    transform: translateX(-50%) !important;
  }
  
  .study-controls {
    top: 10px !important;
    right: 10px !important;
    height: 36px !important;
    gap: 6px !important;
  }
  #study-toggle-btn {
    width: 32px !important;
    height: 32px !important;
    min-width: 32px !important;
    max-width: 32px !important;
    min-height: 32px !important;
    max-height: 32px !important;
    font-size: 14px !important;
  }
  .study-file-upload {
    min-width: 28px !important;
    max-width: 28px !important;
    height: 28px !important;
  }
  .study-file-upload > div {
    height: 28px !important;
    width: 28px !important;
  }
  
  .header-row { height: 56px !important; padding-left: 6px !important; padding-right: 90px !important; gap: 8px !important; }
  .header-title { font-size: 1.1rem !important; }
  #ugm-logo, #ugm-logo .image-container, #ugm-logo > div {
    width: 40px !important; height: 40px !important; min-width: 40px !important; max-width: 40px !important;
  }
  #ugm-logo img { width: 36px !important; height: 36px !important; }
}

/* Android / smaller density screens */
@media (max-width: 400px) {
  .textbox, .textbox textarea {
    font-size: 13px !important;
    min-height: 42px !important;
    padding: 8px 48px 8px 14px !important;
  }
  .textbox textarea::placeholder {
    font-size: 0.72rem !important;
    line-height: 1.3 !important;
  }
}

@media (max-height: 500px) {
  .chatbox { top: 50px !important; bottom: 60px !important; }
  .textbox-container { top: auto !important; bottom: 0px !important; }
  .study-controls { top: 8px !important; right: 8px !important; }
  .header-row { height: 50px !important; }
  .textbox, .textbox textarea { min-height: 40px !important; padding: 10px 45px 10px 14px !important; }
}
"""

# ---------- Blocks layout ----------
with gr.Blocks(title="UGM Chatbot") as demo:
    gr.HTML("""
    <meta name="viewport" content="width=device-width, initial-scale=1, maximum-scale=1, user-scalable=no">
    <style>

    /* Force dark theme globally */
      html, body, .gradio-container, .main, .app, .contain,
      .block, .wrap, .panel, .gap, .form,
      [data-testid="chatbot"], [data-testid="chatbot"] > div {
        background: #0f0f10 !important;
        color: #eee !important;
        color-scheme: dark !important;
      }

      /* Force white text in all bubbles regardless of theme */
      .chatbox .message-bubble, .chatbox .message-bubble *,
      .chatbox .prose, .chatbox .prose * {
        color: #ffffff !important;
      }
      .chatbox .message-row.user-row .message-bubble {
        background: #003d80 !important;
      }
      .chatbox .message-row.bot-row .message-bubble {
        background: #1e1e22 !important;
      }

      /* Remove box around ChatUGM text in header */
      .header-row .svelte-1ed2p3z,
      .header-row .prose,
      .header-row .markdown-text,
      .header-row .block,
      .header-row > div > div {
        background: transparent !important;
        border: none !important;
        box-shadow: none !important;
        padding: 0 !important;
        margin: 0 !important;
        min-height: 0 !important;
      }
    
      #ugm-logo, #ugm-logo > div, #ugm-logo .image-container {
        background: transparent !important; border: none !important;
        padding: 0 !important; margin: 0 !important;
        width: 50px !important; height: 50px !important; min-width: 50px !important;
        box-shadow: none !important;
      }
      .chatbox { left: 0 !important; right: 0 !important; width: 100vw !important;
        max-width: 100vw !important; transform: none !important; margin: 0 !important;
        display: flex !important;}

    /* Show the bin/clear button */
    button[aria-label="Clear"],
    [aria-label="Clear"] {
      display: flex !important;
      visibility: visible !important;
      opacity: 1 !important;
      pointer-events: auto !important;
      width: auto !important;
      height: auto !important;
}
      .chatbox > div { width: 100% !important; max-width: 100% !important; }
      .chatbox .user, .chatbox [data-testid="user"] {
        max-width: 110% !important; width: auto !important; height: auto !important;
        min-height: unset !important; padding: 10px 14px 10px 12px !important; white-space: normal !important;
      }
      .chatbox .user p, .chatbox [data-testid="user"] p { margin: 0 !important; padding: 0 !important; white-space: normal !important; }
      .chatbox .bot, .chatbox .assistant, .chatbox [data-testid="bot"] { padding-left: 12px !important; text-align: left !important; }
      .chatbox .bot p, .chatbox .assistant p, .chatbox [data-testid="bot"] p, .chatbox .prose p {
        margin: 0 !important; padding: 0 !important; text-align: left !important;
      }
      .study-file-upload span, .study-file-upload label, .study-file-upload p,
      .study-file-upload .wrap, .study-file-upload [data-testid="block-info"] {
        display: none !important;
        visibility: hidden !important;
        width: 0 !important;
        height: 0 !important;
        overflow: hidden !important;
      }
      @media (max-width: 768px) {
        .chatbox { bottom: 76px !important; height: auto !important; padding: 2px 2px !important; }
        .textbox-container { top: auto !important; bottom: 0px !important; }
      }
      @media (max-width: 480px) {
        .chatbox { bottom: 72px !important; padding: 2px 2px !important; }
        .textbox-container { top: auto !important; bottom: 0px !important; }
      }
    </style>
    """)
    
    with gr.Column(elem_classes="app-container"):
        with gr.Row(elem_classes="header-row"):
            reset_btn = gr.Button("🎓", elem_id="reset-hat-btn")
            gr.HTML("<span class='header-title'>ChatUGM</span>")

        # Study controls: toggle + file upload (in header area)
        with gr.Row(elem_classes="study-controls"):
            study_toggle = gr.Button("📚", elem_id="study-toggle-btn")
            study_file = gr.File(
                label="📎",
                file_types=[".pdf", ".docx", ".doc", ".pptx", ".ppt", ".xlsx", ".xls", ".csv", ".json", ".txt"],
                file_count="single",
                visible=False,
                elem_classes="study-file-upload"
            )

        chat_display = gr.Chatbot(elem_classes="chatbox", show_label=False)
        state = gr.State({"session_id": str(uuid.uuid4()), "history": []})
        study_mode_on = gr.State(False)

        # Textbox + send button at the bottom
        with gr.Row(elem_classes="textbox-container"):
            txt = gr.Textbox(
                placeholder="Ask me anything about the university...",
                show_label=False, elem_classes="textbox", lines=1, max_lines=8,
                scale=10
            )
            send = gr.Button("", elem_classes="send-btn", scale=1, min_width=48)

        gr.HTML("""
        <script>
          function scrollAllToBottom() {
            document.querySelectorAll('.chatbox, .chatbox > div, .chatbox > div > div, .chatbox .wrap, .chatbox .message-wrap, [data-testid="chatbot"], [data-testid="chatbot"] > div, [data-testid="chatbot"] > div > div, [role="log"]').forEach(function(el) {
              try { el.scrollTop = el.scrollHeight; } catch(e) {}
            });
          }

          // Aggressive: run every 300ms
          setInterval(scrollAllToBottom, 300);

          // Watch for new messages
          function startWatcher() {
            var root = document.querySelector('.chatbox') || document.querySelector('[data-testid="chatbot"]');
            if (!root) { setTimeout(startWatcher, 300); return; }
            new MutationObserver(function() {
              for (var i = 0; i < 10; i++) {
                setTimeout(scrollAllToBottom, i * 100);
              }
            }).observe(root, { childList: true, subtree: true, characterData: true });
          }
          startWatcher();

          // On send
          document.addEventListener('keydown', function(e) {
            if (e.key === 'Enter') {
              for (var i = 1; i <= 20; i++) setTimeout(scrollAllToBottom, i * 100);
            }
          });
          document.addEventListener('click', function() {
            for (var i = 1; i <= 20; i++) setTimeout(scrollAllToBottom, i * 100);
          });

          // Initial
          for (var i = 1; i <= 30; i++) setTimeout(scrollAllToBottom, i * 100);
        </script>
        """)

        def toggle_study_mode(is_on, state):
            state = ensure_session_id(state)
            history = state.get("history", [])
            new_state = not is_on
            
            if new_state:
                msg = "📚 **Study Mode ON!**\n\nUpload your study material (PDF, Word, Excel, PPT, CSV, JSON, TXT) using the 📎 button."
            else:
                bot.study_mode.reset()
                msg = "💬 **Study Mode OFF.** Back to normal chat!"
            
            history.append({"role": "assistant", "content": msg})
            state["history"] = history
            return new_state, history, state, gr.update(visible=new_state)
        
        def on_file_upload(file, is_on, state):
            state = ensure_session_id(state)
            history = state.get("history", [])
            
            if not is_on:
                history.append({"role": "assistant", "content": "⚠️ Please turn on Study Mode first (click the 📚 button)."})
                state["history"] = history
                return history, state
            
            if file is None:
                return history, state
            
            success, msg = bot.study_mode.process_document(file)
            history.append({"role": "assistant", "content": msg})
            state["history"] = history
            return history, state
        
        def clear_chat(state):
            """Permanently clear chat history"""
            new_state = {"session_id": str(uuid.uuid4()), "history": []}
            bot.study_mode.reset()
            return [], new_state, False, gr.update(visible=False)
        
        def chat_send_safe(user_input, is_study_on, state):
            state = ensure_session_id(state)
            session_id = state["session_id"]
            history = state.get("history", [])

            if not user_input.strip():
                if is_study_on and bot.study_mode.document_text:
                    reply = "📚 Type **interact** to ask questions, **assignment** for assignment help, or **summary** for key points!"
                else:
                    reply = "How can I help today? 🙂"
                history.append({"role": "assistant", "content": reply})
                state["history"] = history
                return "", history, state

            # STUDY MODE
            if is_study_on and bot.study_mode.document_text:
                user_lower = user_input.lower().strip()
                
                # Check for assignment mode activation
                if user_lower in ['assignment', 'assignments', 'homework', 'task', 'help with assignment']:
                    reply = bot.study_mode.start_assignment_mode(bot.pipe, bot.tokenizer)
                # Check for interact activation
                elif user_lower in ['interact', 'interactive', 'ask', 'question', 'questions', 'chat', 'talk']:
                    reply = bot.study_mode.start_interact()
                # Check for summary request
                elif user_lower in ['summary', 'summarize', 'summarise', 'sum', 'key points', 'keypoints', 'overview']:
                    reply = bot.study_mode.generate_summary(bot.pipe, bot.tokenizer)
                # Check for exit from modes
                elif user_lower in ['exit', 'quit', 'stop', 'end', 'done']:
                    if bot.study_mode.interact_active:
                        bot.study_mode.interact_active = False
                        reply = "👋 Left interactive mode. Type **interact** to start again, **assignment** for help, or **summary** for overview."
                    elif bot.study_mode.assignment_active:
                        bot.study_mode.assignment_active = False
                        reply = "👋 Left assignment mode. Type **assignment** to start again, **interact** for questions, or **summary** for overview."
                    else:
                        reply = "📚 **Study Options:**\n• Type **interact** for questions\n• Type **assignment** for assignment help\n• Type **summary** for overview"
                # If assignment mode is active, handle the question
                elif bot.study_mode.assignment_active:
                    reply = bot.study_mode.handle_assignment_question(user_input, bot.pipe, bot.tokenizer)
                # If interact mode is active, handle the question
                elif bot.study_mode.interact_active:
                    reply = bot.study_mode.handle_interact(user_input, bot.pipe, bot.tokenizer)
                # Default help message
                else:
                    reply = "📚 **Study Options:**\n• Type **interact** to ask questions about your document\n• Type **assignment** to get help with assignments\n• Type **summary** for key points and overview"
                
                history.append({"role": "user", "content": user_input})
                history.append({"role": "assistant", "content": reply})
                state["history"] = history
                return "", history, state

            # NORMAL CHAT
            reply, tag, score = bot.generate_response(user_input, session_id)
            history.append({"role": "user", "content": user_input})
            history.append({"role": "assistant", "content": reply})
            state["history"] = history
            return "", history, state

        study_toggle.click(
            toggle_study_mode, [study_mode_on, state], [study_mode_on, chat_display, state, study_file]
        ).then(fn=None, js="() => { document.getElementById('study-toggle-btn').classList.toggle('active'); }")

        reset_btn.click(
            clear_chat, [state], [chat_display, state, study_mode_on, study_file]
        ).then(fn=None, js="() => { var sb = document.getElementById('study-toggle-btn'); if(sb) sb.classList.remove('active'); }")
        
        study_file.change(on_file_upload, [study_file, study_mode_on, state], [chat_display, state])
        
        txt.submit(chat_send_safe, [txt, study_mode_on, state], [txt, chat_display, state]).then(
            fn=None, js="() => { setTimeout(() => { document.querySelector('textarea').focus(); }, 100); }"
        )
        send.click(chat_send_safe, [txt, study_mode_on, state], [txt, chat_display, state]).then(
            fn=None, js="() => { setTimeout(() => { document.querySelector('textarea').focus(); }, 100); }"
        )
        
        chat_display.clear(clear_chat, [state], [chat_display, state, study_mode_on, study_file])

if __name__ == "__main__":
    demo.launch(css=css, theme=gr.themes.Default(primary_hue="blue", secondary_hue="red").set(body_background_fill="#0f0f10"))